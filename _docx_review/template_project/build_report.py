from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor


ROOT = Path(r"C:\Users\sriss\Downloads\Deep Learning")
TEMPLATE = Path(r"C:\Users\sriss\Downloads\template_project.docx")
OUTPUT = ROOT / "Brain_Tumor_Classification_Project_Report_Corrected.docx"


def set_paragraph(paragraph, text: str, style: str | None = None, bold: bool = False):
    paragraph.clear()
    if style:
        paragraph.style = style
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def fill_table(table, rows: list[list[str]], widths_dxa: list[int]):
    # The supplied sample tables contain merged header cells. Rebuild the row
    # cells against the existing four-column grid so each heading has its own
    # real cell and cannot overwrite adjacent headings.
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for tr in list(table._tbl.tr_lst):
        table._tbl.remove(tr)
    for row_index, values in enumerate(rows):
        tr = OxmlElement("w:tr")
        tr_pr = OxmlElement("w:trPr")
        if row_index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        tr.append(tr_pr)
        for column_index, value in enumerate(values):
            tc = OxmlElement("w:tc")
            tc_pr = OxmlElement("w:tcPr")
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths_dxa[column_index]))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)
            valign = OxmlElement("w:vAlign")
            valign.set(qn("w:val"), "center")
            tc_pr.append(valign)
            tc.append(tc_pr)
            p = OxmlElement("w:p")
            p_pr = OxmlElement("w:pPr")
            p_style = OxmlElement("w:pStyle")
            p_style.set(qn("w:val"), "tablecolhead" if row_index == 0 else "tablecopy")
            p_pr.append(p_style)
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "center" if row_index == 0 or column_index > 0 else "left")
            p_pr.append(jc)
            p.append(p_pr)
            run = OxmlElement("w:r")
            r_pr = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "000000")
            r_pr.append(color)
            if row_index == 0:
                r_pr.append(OxmlElement("w:b"))
            run.append(r_pr)
            text = OxmlElement("w:t")
            text.text = value
            run.append(text)
            p.append(run)
            tc.append(p)
            tr.append(tc)
        table._tbl.append(tr)


def main():
    doc = Document(TEMPLATE)
    original_paragraphs = list(doc.paragraphs)
    keep = set()

    # Cover title and the three student author blocks.
    set_paragraph(original_paragraphs[0], "Brain Tumor MRI Classification Using Deep Learning", "paper title")
    keep.add(0)
    authors = [
        (3, "Vangara Sridhar (2023BCS0106)", "vangara23bcs106@iiitkottaym.ac.in"),
        (12, "Suragouni Sahithi (2023BCS0070)", "suragou23bcs70@iiitkottaym.ac.in"),
        (20, "Vatte Ramya Sri (2023BCS0208)", "vatte23bcs208@iiitkottaym.ac.in"),
    ]
    for start, name, email in authors:
        values = [name, "Department of Computer Science and Engineering", "", "Indian Institute of Information Technology Kottayam", "", "Kottayam, Kerala, India", email]
        for offset, value in enumerate(values):
            index = start + offset
            set_paragraph(original_paragraphs[index], value)
            keep.add(index)

    # Use the existing author-column section with three columns for the three authors.
    doc.sections[2]._sectPr.xpath("./w:cols")[0].set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num", "3")

    content = {
        41: ("Abstract—This project compares a feedforward artificial neural network, a custom convolutional neural network, and ImageNet-pretrained ResNet50 for four-class brain MRI classification. The local dataset contains 5,600 training images and 1,600 held-out test images across glioma, meningioma, no-tumor, and pituitary classes. ResNet50 achieved 93.69 percent test accuracy and a 93.53 percent macro F1 score, exceeding the custom CNN and ANN baselines. The study demonstrates how spatial feature learning and transfer learning can improve this task, while results remain limited to this dataset and are not clinical validation.", "Abstract"),
        42: ("Keywords—brain MRI, tumor classification, artificial neural network, convolutional neural network, transfer learning", "Keywords"),
        43: ("Introduction", "Heading 1"),
        44: ("Brain MRI contains structural patterns that can help distinguish tumor categories, but manual image review is time-consuming and depends on specialist interpretation. This mini-project studies image-level classification of MRI scans into glioma, meningioma, pituitary tumor, and no-tumor classes. The task is supervised multiclass classification: an image is the input and its dataset label is the target.", "Body Text"),
        45: ("Literature Review", "Heading 1"),
        46: ("Deep learning for image classification", "Heading 2"),
        47: ("Convolutional neural networks learn local image features through convolution and pooling. LeCun and colleagues established convolutional learning as an effective approach for visual pattern recognition [2]. Later image classifiers such as AlexNet, VGG, and Inception developed this approach at greater depth and scale [6]–[8]. Residual networks use identity shortcuts to support optimization of deeper models [3]. Transfer learning research explains how pretrained features can be adapted to new tasks [9], [10]. These ideas motivate the CNN and ResNet50 architectures evaluated here.", "Body Text"),
        48: ("Deep learning is widely studied in medical image analysis [4], [23]. Brain MRI tumor classification has been explored with CNNs and deep networks [20]–[22]. Broader trends include generative adversarial networks [11] and variational autoencoders [12] for generative learning; recurrent networks [13] and Transformers [14] for sequence modeling; Vision Transformers [15] for image recognition; self-supervised contrastive and masked-image pretraining [16], [17]; and Grad-CAM for visual explanations [18]. Data augmentation is surveyed in [19]. These methods provide context; only ANN, custom CNN, and ResNet50 were implemented in this project. Results from different studies are not directly comparable because their datasets and evaluation protocols differ.", "Body Text"),
        49: ("TABLE I. Literature review summary", "table head"),
        51: ("Methodology", "Heading 1"),
        52: ("Dataset and split. The Masoud Nickparvar Brain Tumor MRI Dataset [1] contains four class folders. The local copy has 5,600 images in Training (1,400 per class) and 1,600 images in Testing (400 per class). The archive names the no-tumor folder `notumor`; the project maps it to the canonical label `no_tumor`.", "Body Text"),
        53: ("The Training partition was split once into 4,760 training images and 840 validation images using a stratified 85/15 split with random seed 42. The supplied Testing partition was kept for final evaluation. Images were resized to the model-specific input size; ANN and CNN inputs were scaled to [0,1], while ResNet50 inputs used its ImageNet preprocessing. Random horizontal flips, rotations, zoom, and brightness changes were applied during training only.", "Body Text"),
        58: ("Fig. 1. Project pipeline and evaluated architectures", "figure caption"),
        59: ("Model architectures", "Heading 2"),
        60: ("The ANN baseline receives 64 × 64 RGB images, flattens them to 12,288 values, and applies dense layers of 512 and 256 ReLU units with dropout rate 0.4, followed by a four-unit softmax output. Flattening provides a simple baseline but removes explicit two-dimensional neighborhood structure.", "Body Text"),
        61: ("Custom convolutional neural network", "Heading 2"),
        62: ("The custom CNN receives 128 × 128 RGB images. Three convolutional blocks use 32, 64, and 128 filters with 3 × 3 kernels, batch normalization, ReLU, and 2 × 2 max pooling. Global average pooling and a dense classification head produce four class probabilities.", "Body Text"),
        63: ("ResNet50 transfer learning", "Heading 2"),
        64: ("The ResNet50 model uses a 224 × 224 input and ImageNet-pretrained weights. The backbone is frozen while a new classification head is trained; fine-tuning then unfreezes the upper 30 backbone layers at learning rate 0.00001. Residual blocks compute y = F(x) + x, adding an identity path that helps gradients pass through a deep network [3].", "Body Text"),
        65: ("Objective and architectural rationale", "Heading 2"),
        66: ("All models use a four-class softmax output. For class k, the predicted probability is p(k|x) = exp(z_k) / sum_j exp(z_j). Training minimizes categorical cross-entropy, L = −(1/N) sum_i sum_k y_ik log(p_ik), using the Adam optimizer [5]. Early stopping monitors validation loss and restores the best weights.", "Body Text"),
        67: ("Convolution shares each filter across image locations. A convolutional layer with kernel size K by K, C input channels, and F filters uses approximately K²CF weights, independent of image width and height. A dense layer from D flattened pixels to H units uses DH weights. Weight sharing therefore reduces parameters and encodes local spatial structure. In this implementation, the ANN has about 6.42 million parameters and the custom CNN about 0.13 million; ResNet50 has a larger pretrained backbone but begins with those backbone weights frozen.", "Body Text"),
        68: ("Training used a batch size of 32, Adam with learning rate 0.001, categorical cross-entropy, a maximum of 50 epochs, early-stopping patience 10, and seed 42. ResNet fine-tuning used a lower learning rate and patience 5. Augmentation was restricted to the training subset to avoid altering validation and test measurements.", "Body Text"),
        78: ("Results and Discussion", "Heading 1"),
        79: ("The three saved models were evaluated on the same 1,600-image held-out Testing partition. Accuracy measures the fraction of correct predictions; macro precision, recall, and F1 average the per-class scores equally. The reported metrics are from the project evaluation files.", "Body Text"),
        80: ("Training behavior", "Heading 2"),
        81: ("ResNet50 fine-tuning reached its best checkpoint at epoch 15, with validation loss 0.0609 and validation accuracy 97.26 percent. Training stopped at epoch 20 after validation loss failed to improve for the configured patience period, and the best checkpoint was restored. Early stopping avoids continuing solely to reach the maximum epoch count.", "Body Text"),
        82: ("Performance measures", "Heading 2"),
        83: ("For each class, precision = TP/(TP+FP), recall = TP/(TP+FN), and F1 = 2 × precision × recall/(precision+recall). Macro averages give each of the four classes equal weight. The class-wise metrics help expose errors that overall accuracy can hide.", "Body Text"),
        85: ("Test-set model comparison", "Heading 2"),
        86: ("ResNet50 achieved the strongest aggregate result, with 93.69 percent accuracy and 0.9353 macro F1. The custom CNN reached 70.75 percent accuracy and 0.6808 macro F1, while the ANN reached 59.44 percent and 0.5422. In this experiment, spatial convolution improved on the flattened ANN baseline, and pretrained residual features improved further. These results reflect the full pipeline, so they do not isolate one architectural factor as the cause.", "Body Text"),
        96: ("Fig. 2. ResNet50 confusion matrix on the held-out test set", "figure caption"),
        99: ("Class-wise results", "Heading 2"),
        100: ("ResNet50 recall was 0.7850 for glioma, 0.9650 for meningioma, 1.0000 for no tumor, and 0.9975 for pituitary. Glioma was the weakest class by recall and F1, indicating that some glioma examples were assigned to another category. The ANN also showed low meningioma recall (0.0600), while the CNN showed low glioma recall (0.3725). These differences demonstrate why class-wise reporting is important.", "Body Text"),
        101: ("The results do not establish performance on other hospitals, scanners, patient populations, or image acquisition protocols. The public dataset may contain sampling or labeling limitations; image-level evaluation does not establish patient-level independence. A larger external dataset, patient-wise splitting, confidence intervals, and clinical review would be needed before any real-world application.", "Body Text"),
        103: ("TABLE II. Test set model comparison", "table head"),
        105: ("Interpretation", "Heading 2"),
        106: ("The ANN uses dense connections over flattened pixels, whereas convolutional filters reuse local weights and preserve neighborhood structure. ResNet50 adds pretrained visual features and skip connections. These design choices offer a mathematical and architectural explanation for the observed ranking, but this comparison is not a controlled ablation: input resolution, parameter count, initialization, and training phases also differ.", "Body Text"),
        107: ("TABLE III. ResNet50 per class metrics", "table head"),
        109: ("All reported values are experiment-specific. The model is an educational classification prototype and must not be used to diagnose, exclude, or guide treatment for a brain tumor.", "Body Text"),
        111: ("Conclusion and Future Work", "Heading 1"),
        112: ("This project trained and compared an ANN, a custom CNN, and ResNet50 transfer learning for four-class brain MRI classification. ResNet50 produced the best result on the supplied held-out test set, with 93.69 percent accuracy and 0.9353 macro F1. The CNN improved on the ANN baseline, and class-wise results identified glioma recall as the main weakness of ResNet50. Future work should test patient-wise and external validation, quantify uncertainty, inspect model attention with explainability methods such as Grad-CAM, and assess additional datasets and imaging protocols. The current system is for coursework and research demonstration only.", "Body Text"),
        114: ("References", "Heading 5"),
        115: ("[1] M. Nickparvar, Brain Tumor MRI Dataset, Kaggle. [Online]. Available: https://www.kaggle.com/datasets/masoudnickparvar/brain-tumor-mri-dataset", "references"),
        116: ("[2] Y. LeCun, L. Bottou, Y. Bengio, and P. Haffner, “Gradient-based learning applied to document recognition,” Proceedings of the IEEE, vol. 86, no. 11, pp. 2278–2324, 1998.", "references"),
        117: ("[3] K. He, X. Zhang, S. Ren, and J. Sun, “Deep residual learning for image recognition,” in Proc. IEEE Conference on Computer Vision and Pattern Recognition, pp. 770–778, 2016.", "references"),
        118: ("[4] G. Litjens et al., “A survey on deep learning in medical image analysis,” Medical Image Analysis, vol. 42, pp. 60–88, 2017.", "references"),
        119: ("[5] D. P. Kingma and J. Ba, “Adam: A method for stochastic optimization,” in Proc. International Conference on Learning Representations, 2015.", "references"),
    }
    for index, (text, style) in content.items():
        set_paragraph(original_paragraphs[index], text, style)
        keep.add(index)

    more_references = [
        "[6] A. Krizhevsky, I. Sutskever, and G. E. Hinton, “ImageNet classification with deep convolutional neural networks,” in Advances in Neural Information Processing Systems, vol. 25, 2012, pp. 1097–1105.",
        "[7] K. Simonyan and A. Zisserman, “Very deep convolutional networks for large-scale image recognition,” in Proc. International Conference on Learning Representations, 2015.",
        "[8] C. Szegedy et al., “Going deeper with convolutions,” in Proc. IEEE Conference on Computer Vision and Pattern Recognition, 2015, pp. 1–9.",
        "[9] S. J. Pan and Q. Yang, “A survey on transfer learning,” IEEE Transactions on Knowledge and Data Engineering, vol. 22, no. 10, pp. 1345–1359, 2010.",
        "[10] J. Yosinski, J. Clune, Y. Bengio, and H. Lipson, “How transferable are features in deep neural networks?” in Advances in Neural Information Processing Systems, vol. 27, 2014.",
        "[11] I. Goodfellow et al., “Generative adversarial nets,” in Advances in Neural Information Processing Systems, vol. 27, 2014, pp. 2672–2680.",
        "[12] D. P. Kingma and M. Welling, “Auto-encoding variational Bayes,” in Proc. International Conference on Learning Representations, 2014.",
        "[13] S. Hochreiter and J. Schmidhuber, “Long short-term memory,” Neural Computation, vol. 9, no. 8, pp. 1735–1780, 1997.",
        "[14] A. Vaswani et al., “Attention is all you need,” in Advances in Neural Information Processing Systems, vol. 30, 2017, pp. 5998–6008.",
        "[15] A. Dosovitskiy et al., “An image is worth 16×16 words: Transformers for image recognition at scale,” in Proc. International Conference on Learning Representations, 2021.",
        "[16] T. Chen et al., “A simple framework for contrastive learning of visual representations,” in Proc. International Conference on Machine Learning, PMLR, vol. 119, 2020, pp. 1597–1607.",
        "[17] K. He et al., “Masked autoencoders are scalable vision learners,” in Proc. IEEE/CVF Conference on Computer Vision and Pattern Recognition, 2022, pp. 16000–16009.",
        "[18] R. R. Selvaraju et al., “Grad-CAM: Visual explanations from deep networks via gradient-based localization,” in Proc. IEEE International Conference on Computer Vision, 2017, pp. 618–626.",
        "[19] C. Shorten and T. M. Khoshgoftaar, “A survey on image data augmentation for deep learning,” Journal of Big Data, vol. 6, art. 60, 2019.",
        "[20] N. Abiwinanda et al., “Brain tumor classification using convolutional neural network,” in Proc. World Congress on Medical Physics and Biomedical Engineering, 2018, pp. 183–189.",
        "[21] H. H. Sultan, N. M. Salem, and W. Al-Atabany, “Multi-classification of brain tumor images using deep neural network,” IEEE Access, vol. 7, pp. 69215–69225, 2019.",
        "[22] J. Cheng et al., “Enhanced performance of brain tumor classification via tumor region augmentation and partition,” PLOS ONE, vol. 10, no. 10, e0140381, 2015.",
        "[23] S. Bauer, R. Wiest, L.-P. Nolte, and M. Reyes, “A survey of MRI-based medical image analysis for brain tumor studies,” Physics in Medicine and Biology, vol. 58, no. 13, pp. R97–R129, 2013.",
    ]
    section_end = original_paragraphs[131]
    for reference in more_references:
        paragraph = section_end.insert_paragraph_before(reference, style="references")
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

    # Retain section-boundary paragraphs and figure anchors; remove the template's
    # instructional prose and unused fourth author column.
    keep.update({2, 37, 40, 57, 95, 131})
    for index, paragraph in enumerate(original_paragraphs):
        if index not in keep:
            paragraph._p.getparent().remove(paragraph._p)

    # Reuse the three template tables but size them for report content.
    fill_table(doc.tables[0], [
        ["Reference", "Method", "Contribution", "Relevance"],
        ["LeCun et al. [2]", "CNN", "Learned local visual features", "Basis for image CNNs"],
        ["He et al. [3]", "ResNet", "Identity shortcut connections", "Supports deep feature learning"],
        ["Litjens et al. [4]", "Medical imaging survey", "Reviews deep learning methods", "Context for evaluation"],
    ], [850, 1050, 1650, 1310])
    fill_table(doc.tables[1], [
        ["Model", "Accuracy", "Macro F1", "Interpretation"],
        ["ANN", "0.5944", "0.5422", "Flattened baseline"],
        ["Custom CNN", "0.7075", "0.6808", "Learns spatial features"],
        ["ResNet50", "0.9369", "0.9353", "ImageNet transfer learning"],
    ], [1200, 950, 950, 1760])
    fill_table(doc.tables[2], [
        ["Class", "Precision", "Recall", "F1 score"],
        ["Glioma", "0.9937", "0.7850", "0.8771"],
        ["Meningioma", "0.8616", "0.9650", "0.9104"],
        ["No tumor", "0.9479", "1.0000", "0.9732"],
        ["Pituitary", "0.9638", "0.9975", "0.9803"],
    ], [1350, 1170, 1170, 1170])

    # Replace the two template placeholder images while preserving their anchors.
    image_paths = [ROOT / "_docx_review" / "architecture.png", ROOT / "results" / "resnet" / "confusion_matrix.png"]
    for shape, image_path in zip(doc.inline_shapes, image_paths):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        doc.part.related_parts[rid]._blob = image_path.read_bytes()

    # Keep all supplied text normally readable in the black text expected by this paper template.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

    props = doc.core_properties
    props.title = "Brain Tumor MRI Classification Using Deep Learning"
    props.subject = "Mini project report comparing ANN, CNN, and ResNet50"
    props.author = "Vangara Sridhar; Suragouni Sahithi; Vatte Ramya Sri"
    props.keywords = "brain MRI, classification, deep learning, CNN, ResNet50"
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
