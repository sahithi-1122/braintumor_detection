from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import RGBColor


root = Path(r"C:\Users\sriss\Downloads\Deep Learning")
source = root / "Brain_Tumor_Classification_Project_Report_Corrected.docx"
output = root / "Brain_Tumor_Classification_Project_Report_Final.docx"
url = "https://github.com/sahithi-1122/braintumor_detection"

doc = Document(source)
section_end = next(
    paragraph
    for paragraph in reversed(doc.paragraphs)
    if paragraph._p.pPr is not None and paragraph._p.pPr.sectPr is not None
)

heading = section_end.insert_paragraph_before("Project Repository", style="Heading 5")
for run in heading.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

paragraph = section_end.insert_paragraph_before(style="references")
prefix = paragraph.add_run("GitHub repository: ")
prefix.font.color.rgb = RGBColor(0, 0, 0)

relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
hyperlink = OxmlElement("w:hyperlink")
hyperlink.set(qn("r:id"), relationship_id)
run = OxmlElement("w:r")
properties = OxmlElement("w:rPr")
color = OxmlElement("w:color")
color.set(qn("w:val"), "0563C1")
properties.append(color)
underline = OxmlElement("w:u")
underline.set(qn("w:val"), "single")
properties.append(underline)
run.append(properties)
text = OxmlElement("w:t")
text.text = url
run.append(text)
hyperlink.append(run)
paragraph._p.append(hyperlink)

doc.save(output)
print(output)
